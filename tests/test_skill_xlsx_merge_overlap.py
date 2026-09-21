"""Issue #3280: an overlapping merge target was written silently.

``edit_xlsx``'s ``merge_cells`` op and ``create_xlsx``'s ``merged`` spec handed
the range to ``ws.merge_cells`` with no overlap check. openpyxl accepts a
range that intersects an existing merge and writes a workbook with
intersecting ``mergeCell`` entries -- invalid content that Excel reports as
corrupt and repairs on open -- while the run reported ``{"applied": 1}``.

A malformed range already failed loudly and left nothing written (#1993). An
overlapping one now fails the same way, before anything is written, with a
message naming both ranges. The malformed path is pinned unchanged.
"""

from __future__ import annotations

import json
import sys
from pathlib import Path
from typing import Any

import pytest
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
SCRIPTS = ROOT / "src" / "agentos" / "skills" / "bundled" / "xlsx" / "scripts"


def _scripts() -> tuple[Any, Any]:
    sys.path.insert(0, str(SCRIPTS))
    try:
        import create_xlsx  # type: ignore[import-not-found]
        import edit_xlsx  # type: ignore[import-not-found]
    finally:
        sys.path.pop(0)
    return create_xlsx, edit_xlsx


def _merged(path: Path, sheet: str = "Form") -> list[str]:
    return sorted(r.coord for r in load_workbook(str(path))[sheet].merged_cells.ranges)


def _source(tmp_path: Path, merged: list[str]) -> Path:
    """The issue's workbook: one sheet, ``merged`` as given."""
    create_xlsx, _ = _scripts()
    src = tmp_path / "in.xlsx"
    create_xlsx.build(
        {"sheets": [{"name": "Form", "rows": [["a", "b", "c", "d"]], "merged": merged}]}
    ).save(str(src))
    return src


def _edit_cli(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path, src: Path, ops: list[dict[str, Any]]
) -> tuple[int, Path]:
    """Drive ``edit_xlsx.main`` the way a caller does; returns (exit, out path)."""
    _, edit_xlsx = _scripts()
    ops_path = tmp_path / "ops.json"
    ops_path.write_text(json.dumps(ops), encoding="utf-8")
    out = tmp_path / "out.xlsx"
    monkeypatch.setattr(sys, "argv", ["edit_xlsx.py", str(src), str(ops_path), "--out", str(out)])
    return edit_xlsx.main(), out


def _merge_op(rng: str, sheet: str = "Form") -> dict[str, Any]:
    return {"op": "merge_cells", "sheet": sheet, "range": rng}


# Targets that intersect an existing A1:B1 merge, in every way a range can.
OVERLAPS = [
    pytest.param("B1:C1", id="shares_an_edge_cell"),
    pytest.param("A1:B1", id="identical"),
    pytest.param("A1:D1", id="encloses"),
    pytest.param("B1", id="single_cell_inside"),
    pytest.param("B1:B3", id="crosses_vertically"),
    pytest.param("A1:A1", id="single_cell_at_the_corner"),
]

# Targets that do not touch A1:B1.
DISJOINT = [
    pytest.param("C1:D1", id="adjacent_on_the_right"),
    pytest.param("A2:B2", id="adjacent_below"),
    pytest.param("C2:D3", id="elsewhere"),
]


# ── edit_xlsx: the issue's steps ────────────────────────────────────────────


def test_the_issues_exact_reproduction(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    """Create with ``"merged": ["A1:B1"]``, apply ``merge_cells B1:C1``: used to
    exit 0 with ``{"applied": 1}`` and write ``['A1:B1', 'B1:C1']``."""
    src = _source(tmp_path, ["A1:B1"])

    with pytest.raises(ValueError, match=r"B1:C1.*overlaps.*A1:B1"):
        _edit_cli(monkeypatch, tmp_path, src, [_merge_op("B1:C1")])

    assert not (tmp_path / "out.xlsx").exists(), "an invalid workbook was written"
    assert capsys.readouterr().out == "", "no success report either"
    assert _merged(src) == ["A1:B1"], "the input is untouched"


@pytest.mark.parametrize("rng", OVERLAPS)
def test_every_overlap_shape_is_refused_before_writing(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path, rng: str
) -> None:
    src = _source(tmp_path, ["A1:B1"])

    with pytest.raises(ValueError) as excinfo:
        _edit_cli(monkeypatch, tmp_path, src, [_merge_op(rng)])

    message = str(excinfo.value)
    assert rng in message and "A1:B1" in message and "'Form'" in message
    assert not (tmp_path / "out.xlsx").exists()


@pytest.mark.parametrize("rng", DISJOINT)
def test_a_disjoint_merge_still_applies(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path, capsys: pytest.CaptureFixture[str], rng: str
) -> None:
    src = _source(tmp_path, ["A1:B1"])

    rc, out = _edit_cli(monkeypatch, tmp_path, src, [_merge_op(rng)])

    assert rc == 0
    assert json.loads(capsys.readouterr().out) == {"applied": 1}
    assert _merged(out) == sorted(["A1:B1", rng])


def test_an_overlap_within_the_same_ops_batch_is_refused(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    """The second op must see the merge the first one made."""
    src = _source(tmp_path, [])

    with pytest.raises(ValueError, match=r"B1:C1.*overlaps.*A1:B1"):
        _edit_cli(monkeypatch, tmp_path, src, [_merge_op("A1:B1"), _merge_op("B1:C1")])

    assert not (tmp_path / "out.xlsx").exists()


def test_earlier_ops_are_not_written_when_a_later_merge_overlaps(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    """All-or-nothing, as for a malformed range: a set_cell before the bad
    merge does not reach the output either."""
    src = _source(tmp_path, ["A1:B1"])
    ops = [
        {"op": "set_cell", "sheet": "Form", "row": 2, "col": 1, "value": "edited"},
        _merge_op("B1:C1"),
    ]

    with pytest.raises(ValueError):
        _edit_cli(monkeypatch, tmp_path, src, ops)

    assert not (tmp_path / "out.xlsx").exists()
    assert load_workbook(str(src))["Form"]["A2"].value is None


def test_the_check_is_per_sheet(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    """A merge on another sheet is not an overlap."""
    create_xlsx, _ = _scripts()
    src = tmp_path / "in.xlsx"
    create_xlsx.build(
        {
            "sheets": [
                {"name": "Form", "rows": [["a", "b"]], "merged": ["A1:B1"]},
                {"name": "Other", "rows": [["a", "b"]]},
            ]
        }
    ).save(str(src))

    rc, out = _edit_cli(monkeypatch, tmp_path, src, [_merge_op("A1:B1", sheet="Other")])

    assert rc == 0
    assert _merged(out, "Other") == ["A1:B1"]


@pytest.mark.parametrize("rng", ["not-a-range", "A1:B"])
def test_a_malformed_range_fails_exactly_as_before(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path, rng: str
) -> None:
    """Pinned unchanged: same exception type and message as openpyxl's own."""
    src = _source(tmp_path, ["A1:B1"])

    with pytest.raises(ValueError, match="not a valid coordinate or range"):
        _edit_cli(monkeypatch, tmp_path, src, [_merge_op(rng)])

    assert not (tmp_path / "out.xlsx").exists()


def test_a_merge_op_naming_a_missing_sheet_is_still_skipped_not_refused(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    """Scope guard: the existing skip for an unknown sheet is not turned into
    an error by the overlap check."""
    src = _source(tmp_path, ["A1:B1"])

    rc, out = _edit_cli(monkeypatch, tmp_path, src, [_merge_op("B1:C1", sheet="Nope")])

    assert rc == 0
    assert json.loads(capsys.readouterr().out) == {"applied": 0}
    assert _merged(out) == ["A1:B1"]


# ── create_xlsx: the merged spec ────────────────────────────────────────────


def test_create_refuses_an_overlapping_spec(tmp_path: Path) -> None:
    create_xlsx, _ = _scripts()

    with pytest.raises(ValueError, match=r"B1:C1.*overlaps.*A1:B1"):
        create_xlsx.build({"sheets": [{"name": "Form", "merged": ["A1:B1", "B1:C1"]}]})


@pytest.mark.parametrize(
    "merged",
    [
        pytest.param(["A1:B1", {"range": "B1:C1"}], id="string_then_object"),
        pytest.param([{"range": "A1:B1"}, "B1:C1"], id="object_then_string"),
    ],
)
def test_create_checks_both_merged_spellings(merged: list[object]) -> None:
    """Strings and ``{"range": ...}`` objects share the one check."""
    create_xlsx, _ = _scripts()

    with pytest.raises(ValueError, match=r"B1:C1.*overlaps.*A1:B1"):
        create_xlsx.build({"sheets": [{"name": "Form", "merged": merged}]})


def test_create_still_accepts_disjoint_merges(tmp_path: Path) -> None:
    create_xlsx, _ = _scripts()
    out = tmp_path / "c.xlsx"

    create_xlsx.build(
        {"sheets": [{"name": "Form", "merged": ["A1:B1", {"range": "C1:D1"}, "A2:B3"]}]}
    ).save(str(out))

    assert _merged(out) == ["A1:B1", "A2:B3", "C1:D1"]


def test_create_main_writes_nothing_on_an_overlap(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    """Through ``main``: the refusal happens before ``wb.save``."""
    create_xlsx, _ = _scripts()
    spec = tmp_path / "spec.json"
    spec.write_text(
        json.dumps({"sheets": [{"name": "Form", "merged": ["A1:B1", "B1:C1"]}]}),
        encoding="utf-8",
    )
    out = tmp_path / "c.xlsx"
    monkeypatch.setattr(sys, "argv", ["create_xlsx.py", str(spec), "--out", str(out)])

    with pytest.raises(ValueError, match="overlaps"):
        create_xlsx.main()

    assert not out.exists()


@pytest.mark.parametrize("merged", [["not-a-range"], [{"range": "not-a-range"}]])
def test_create_malformed_ranges_fail_exactly_as_before(merged: list[object]) -> None:
    create_xlsx, _ = _scripts()

    with pytest.raises(ValueError, match="not a valid coordinate or range"):
        create_xlsx.build({"sheets": [{"merged": merged}]})


# ── the produced file is what Excel would accept ───────────────────────────


def test_a_workbook_that_passes_the_check_has_no_intersecting_merges(tmp_path: Path) -> None:
    """The property behind the fix: every pair of merges in the output is
    disjoint, which is what Excel requires of ``mergeCells``."""
    from openpyxl.worksheet.cell_range import CellRange

    create_xlsx, _ = _scripts()
    out = tmp_path / "c.xlsx"
    create_xlsx.build(
        {"sheets": [{"name": "Form", "merged": ["A1:C1", "D1:D3", "A2:B2", "C2:C4"]}]}
    ).save(str(out))

    ranges = [CellRange(r) for r in _merged(out)]
    for i, left in enumerate(ranges):
        for right in ranges[i + 1 :]:
            assert left.isdisjoint(right), (left.coord, right.coord)
