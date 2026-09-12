from __future__ import annotations

import io
import json
import sys
from pathlib import Path

import pytest

# Ensure the pptx skill script is importable
pptx_script_dir = Path(__file__).resolve().parents[2] / "src/agentos/skills/bundled/pptx/scripts"
if str(pptx_script_dir) not in sys.path:
    sys.path.insert(0, str(pptx_script_dir))

import extract_text  # noqa: E402


def test_pptx_extract_nested_group_shapes(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    pptx = pytest.importorskip("pptx")
    from pptx.util import Inches

    prs = pptx.Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # Top-level textbox
    tb0 = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(1))
    tb0.text_frame.text = "Top Level Text"

    # Top-level table
    table_shape = slide.shapes.add_table(1, 2, Inches(1), Inches(5), Inches(2), Inches(1))
    table_shape.table.cell(0, 0).text = "Cell A"
    table_shape.table.cell(0, 1).text = "Cell B"

    # Group 1 (depth 1)
    g1 = slide.shapes.add_group_shape()
    tb1 = g1.shapes.add_textbox(Inches(1), Inches(2), Inches(2), Inches(1))
    tb1.text_frame.text = "Depth 1 Text"

    # Group 2 nested inside Group 1 (depth 2)
    g2 = g1.shapes.add_group_shape()
    tb2 = g2.shapes.add_textbox(Inches(1), Inches(3), Inches(2), Inches(1))
    tb2.text_frame.text = "Depth 2 Text"

    # Group 3 nested inside Group 2 (depth 3)
    g3 = g2.shapes.add_group_shape()
    tb3 = g3.shapes.add_textbox(Inches(1), Inches(4), Inches(2), Inches(1))
    tb3.text_frame.text = "Depth 3 Text"

    # Verify _slide_text extracts all nested levels
    extracted = extract_text._slide_text(slide)
    assert "Top Level Text" in extracted
    assert "Depth 1 Text" in extracted
    assert "Depth 2 Text" in extracted
    assert "Depth 3 Text" in extracted
    assert "Cell A | Cell B" in extracted

    # Save to file and verify CLI main() with --json
    deck_path = tmp_path / "nested_groups.pptx"
    prs.save(str(deck_path))

    stdout_capture = io.StringIO()
    monkeypatch.setattr(sys, "stdout", stdout_capture)

    rc = extract_text.main([str(deck_path), "--json"])
    assert rc == 0

    records = json.loads(stdout_capture.getvalue())
    assert len(records) == 1
    slide_record = records[0]
    assert slide_record["slide"] == 1
    assert "Top Level Text" in slide_record["text"]
    assert "Depth 1 Text" in slide_record["text"]
    assert "Depth 2 Text" in slide_record["text"]
    assert "Depth 3 Text" in slide_record["text"]
    assert "Cell A | Cell B" in slide_record["text"]
