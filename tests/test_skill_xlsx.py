"""xlsx skill — load, eligibility, and create→inspect→edit round-trip."""

from __future__ import annotations

import sys
from pathlib import Path

import pytest

from agentos.skills.eligibility import EligibilityContext, check_eligibility
from agentos.skills.loader import SkillLoader

ROOT = Path(__file__).resolve().parents[1]
BUNDLED = ROOT / "src" / "agentos" / "skills" / "bundled"
SCRIPTS = BUNDLED / "xlsx" / "scripts"


def _spec() -> object:
    return SkillLoader(bundled_dir=BUNDLED).get_by_name("xlsx")


def test_skill_loads() -> None:
    spec = _spec()
    assert spec is not None
    assert spec.name == "xlsx"
    assert spec.metadata is not None


def test_eligibility_with_python(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(
        "agentos.skills.eligibility.shutil.which",
        lambda name: "/usr/bin/python3" if name in {"python", "python3"} else None,
    )
    spec = _spec()
    assert spec is not None
    assert check_eligibility(spec, EligibilityContext.auto())


@pytest.mark.parametrize(
    "merged",
    [
        pytest.param([{"range": "A1:B1"}], id="dictionary"),
        pytest.param(["A1:B1"], id="string"),
    ],
)
def test_round_trip_with_formula_and_merge(tmp_path: Path, merged: list[object]) -> None:
    sys.path.insert(0, str(SCRIPTS))
    try:
        import create_xlsx  # type: ignore[import-not-found]
        import edit_xlsx  # type: ignore[import-not-found]
        import inspect_xlsx  # type: ignore[import-not-found]
    finally:
        sys.path.pop(0)

    spec = {
        "sheets": [
            {
                "name": "Sales",
                "rows": [
                    ["Region", "Revenue"],
                    ["NA", 1_200_000],
                    ["EU", 850_000],
                    ["Total", "=SUM(B2:B3)"],
                ],
                "merged": merged,
                "freeze": "A2",
            }
        ]
    }
    src = tmp_path / "book.xlsx"
    create_xlsx.build(spec).save(str(src))
    assert src.exists()

    inspected = inspect_xlsx.inspect(src, data_only=False)
    sheet = next(s for s in inspected["sheets"] if s["name"] == "Sales")
    assert sheet["max_row"] == 4
    assert sheet["max_col"] == 2
    assert sheet["merged"] == ["A1:B1"]
    assert sheet["freeze"] == "A2"

    last_row = sheet["rows"][3]
    assert last_row[1]["type"] == "f"
    assert last_row[1]["value"] == "=SUM(B2:B3)"

    from openpyxl import load_workbook

    wb = load_workbook(str(src))
    edit_xlsx.apply_ops(
        wb,
        [
            {"op": "set_cell", "sheet": "Sales", "row": 2, "col": 1, "value": "Americas"},
            {"op": "rename_sheet", "old": "Sales", "new": "Q3"},
        ],
    )
    out = tmp_path / "out.xlsx"
    wb.save(str(out))

    re_inspected = inspect_xlsx.inspect(out, data_only=False)
    sheet_q3 = next(s for s in re_inspected["sheets"] if s["name"] == "Q3")
    assert sheet_q3["rows"][1][0]["value"] == "Americas"


@pytest.mark.parametrize(
    ("merged", "expected"),
    [
        pytest.param(["A1:B1", {"range": "A3:B3"}], ["A1:B1", "A3:B3"], id="mixed"),
        pytest.param([], [], id="empty"),
        pytest.param(None, [], id="null"),
        pytest.param([None, 42, {}, {"other": "A1:B1"}], [], id="unsupported-entries"),
    ],
)
def test_create_merge_formats(tmp_path: Path, merged: object, expected: list[str]) -> None:
    from agentos.skills.bundled.xlsx.scripts.create_xlsx import build
    from agentos.skills.bundled.xlsx.scripts.inspect_xlsx import inspect

    path = tmp_path / "merges.xlsx"
    wb = build({"sheets": [{"name": "Merges", "merged": merged}, {"name": "Plain"}]})
    wb.save(path)
    wb.close()

    sheets = inspect(path, data_only=False)["sheets"]
    assert sheets[0]["merged"] == expected
    assert sheets[1]["merged"] == []


@pytest.mark.parametrize("merged", [["not-a-range"], [{"range": "not-a-range"}]])
def test_create_rejects_invalid_merge_ranges(merged: list[object]) -> None:
    from agentos.skills.bundled.xlsx.scripts.create_xlsx import build

    with pytest.raises(ValueError, match="not a valid coordinate or range"):
        build({"sheets": [{"merged": merged}]})


def test_create_accepts_inspected_merge_ranges(tmp_path: Path) -> None:
    from openpyxl import Workbook

    from agentos.skills.bundled.xlsx.scripts.create_xlsx import build
    from agentos.skills.bundled.xlsx.scripts.inspect_xlsx import inspect

    original = Workbook()
    ws = original.active
    assert ws is not None
    ws.merge_cells("A1:C1")
    ws.merge_cells("B3:B5")
    source = tmp_path / "source.xlsx"
    original.save(source)
    original.close()
    merges = inspect(source, data_only=False)["sheets"][0]["merged"]
    assert merges == ["A1:C1", "B3:B5"]

    # Reuse only merge metadata: inspector rows have a different schema.
    rebuilt = build({"sheets": [{"merged": merges}]})
    target = tmp_path / "rebuilt.xlsx"
    rebuilt.save(target)
    rebuilt.close()
    assert inspect(target, data_only=False)["sheets"][0]["merged"] == merges


def test_text_escapes_formula(tmp_path: Path) -> None:
    sys.path.insert(0, str(SCRIPTS))
    try:
        import create_xlsx  # type: ignore[import-not-found]
        import edit_xlsx  # type: ignore[import-not-found]
        import inspect_xlsx  # type: ignore[import-not-found]
    finally:
        sys.path.pop(0)

    src = tmp_path / "book.xlsx"
    create_xlsx.build({"sheets": [{"name": "S", "rows": [["a"]]}]}).save(str(src))

    from openpyxl import load_workbook

    wb = load_workbook(str(src))
    edit_xlsx.apply_ops(
        wb,
        [
            {
                "op": "set_cell",
                "sheet": "S",
                "row": 2,
                "col": 1,
                "value": "=hello",
                "as_text": True,
            },
        ],
    )
    out = tmp_path / "out.xlsx"
    wb.save(str(out))

    inspected = inspect_xlsx.inspect(out, data_only=False)
    sheet = inspected["sheets"][0]
    cell_value = sheet["rows"][1][0]["value"]
    assert isinstance(cell_value, str)
    assert cell_value.lstrip("'") == "=hello"


def test_inspect_xlsx_creates_parent_directory(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    sys.path.insert(0, str(SCRIPTS))
    try:
        import create_xlsx  # type: ignore[import-not-found]
        import inspect_xlsx  # type: ignore[import-not-found]
    finally:
        sys.path.pop(0)

    src = tmp_path / "book.xlsx"
    create_xlsx.build({"sheets": [{"name": "S", "rows": [["a"]]}]}).save(str(src))

    out = tmp_path / "nested" / "dir" / "out.json"
    monkeypatch.setattr(sys, "argv", ["inspect_xlsx.py", str(src), "--out", str(out)])
    assert inspect_xlsx.main() == 0
    assert out.is_file()
